"""Read the PS2 tab of the Texture Packs Archive spreadsheet.

Columns are found by their header text, not by position, so a column added or moved is either
read correctly or reported, never read as the wrong field. The mirror column is never read: those
links are the archive maintainer's own hosting and stay out of this project."""
from __future__ import annotations

import io
import re
import urllib.request
from dataclasses import dataclass, field

SHEET_ID = "1sif8FeRGJRbytK8wFRXgF6Hke9V6GUFs"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
TAB = "PS2"

# header text (lower case, spaces collapsed) -> field
HEADERS = {
    "game title": "title",
    "download": "download",
    "region": "region",
    "serial number": "serial",
    "size": "size",
    "restriction": "restriction",
    "status": "status",
    "type": "type",
    "pack author": "author",
    "tip the author": "tip",
    "author's socials": "socials",
}
REQUIRED = ("title", "download", "author", "restriction")


@dataclass
class Row:
    row: int
    title: str
    download: str | None
    region: str | None = None
    serial: str | None = None
    size: str | None = None
    restriction: str | None = None
    status: str | None = None
    type: str | None = None
    author: str | None = None
    author_page: str | None = None
    tip: str | None = None
    socials: str | None = None
    extra: dict = field(default_factory=dict)

    @property
    def paid(self) -> bool:
        return (self.restriction or "").strip().lower() == "paywall"

    def size_bytes(self) -> int | None:
        return parse_size(self.size)


def parse_size(text) -> int | None:
    """'1.11 GB', '5,03 GB', '459 MB' -> bytes; anything else -> None."""
    m = re.match(r"\s*([\d.,]+)\s*([KMGT])i?B\b", str(text or ""), re.I)
    if not m:
        return None
    number = m.group(1).replace(",", ".")
    if number.count(".") > 1:
        return None
    return int(float(number) * 1000 ** "KMGT".index(m.group(2).upper()) * 1000)


def fetch(url: str = EXPORT_URL) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "ps2ktxpak-scraper"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def _text(cell) -> str | None:
    v = cell.value
    if v is None:
        return None
    v = str(v).strip()
    return None if v in ("", "N/A", "n/a") else v


def _link(cell) -> str | None:
    target = cell.hyperlink.target if cell.hyperlink else None
    return target.strip() if target else None


def parse(data: bytes, tab: str = TAB) -> list[Row]:
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data))[tab]
    header = [re.sub(r"\s+", " ", str(c.value or "")).strip().lower() for c in ws[1]]
    cols = {}
    for i, h in enumerate(header):
        if h in HEADERS:
            cols[HEADERS[h]] = i
    missing = [f for f in REQUIRED if f not in cols]
    if missing:
        raise ValueError(f"spreadsheet tab {tab!r} has no column for {missing}; headers are {header}")

    rows = []
    for n, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        get = lambda f: cells[cols[f]] if f in cols and cols[f] < len(cells) else None
        title = _text(get("title"))
        if not title:
            continue
        r = Row(row=n, title=title, download=_link(get("download")))
        for f in ("region", "serial", "size", "restriction", "status", "type"):
            c = get(f)
            setattr(r, f, _text(c) if c is not None else None)
        if (c := get("author")) is not None:
            r.author, r.author_page = _text(c), _link(c)
        for f in ("tip", "socials"):
            if (c := get(f)) is not None:
                setattr(r, f, _link(c))
        rows.append(r)
    return rows
